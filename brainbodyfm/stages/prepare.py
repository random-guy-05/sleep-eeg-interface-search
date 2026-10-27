"""Build the STAGES cohort from immutable source files; never fit a predictor."""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path
import zipfile

from openpyxl import load_workbook
import pandas as pd

BASE = Path(__file__).resolve().parent


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def excel_rows(source, *, sheet=None, max_col=None):
    w = load_workbook(source, read_only=True, data_only=True)
    s = w[sheet] if sheet else w.active
    h = next(s.iter_rows(max_row=1, values_only=True))
    try:
        for rownum, values in enumerate(s.iter_rows(min_row=2, max_col=max_col, values_only=True), 2):
            if any(v is not None for v in values):
                yield rownum, dict(zip(h, values))
    finally:
        w.close()


def num(x):
    try:
        value = float(x)
        return value if math.isfinite(value) else None
    except (ValueError, TypeError):
        return None


def text(x):
    return '' if x is None else str(x).strip()


def day(x):
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    value = pd.to_datetime(x, errors='coerce')
    return None if pd.isna(value) else value.date()


def sex_code(x):
    return {'F': 0, 'FEMALE': 0, 'M': 1, 'MALE': 1}.get(text(x).upper())


def truth(x):
    return text(x).lower() in {'t', 'true', 'yes', '1'}


def total(items, maximum=3):
    values = [num(x) for x in items]
    return sum(values) if all(v is not None and v.is_integer() and 0 <= v <= maximum for v in values) else None


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--base', type=Path, default=BASE)
    args = parser.parse_args()
    base = args.base
    src, out = base / 'data/source', base / 'data'
    protocol_path = base / 'protocol_v1.json'
    cfg = json.loads(protocol_path.read_text())
    protocol_paths = [protocol_path]
    amendment = base / 'amendment_01.json'
    if amendment.exists():
        cfg.update(json.loads(amendment.read_text())['overrides'])
        protocol_paths.append(amendment)
    protocol_sha = hashlib.sha256(b'\n'.join(p.read_bytes() for p in protocol_paths)).hexdigest()
    # Input hashes are checked before every build.
    for item in json.loads((src / 'SHA256SUMS.json').read_text()):
        p = src / item['file']
        if p.stat().st_size != item['bytes'] or sha(p) != item['sha256']:
            raise ValueError(f'Source snapshot changed: {p.name}')
    inventory = json.loads((out / 'psg_drive_inventory.json').read_text())
    local_inventory = (out / 'psg_local_inventory.json')
    local = json.loads(local_inventory.read_text()) if local_inventory.exists() else None
    edfs = defaultdict(list)
    drive_by_name = {}
    for folder in inventory['folders']:
        for f in folder['files']:
            if f['name'].lower().endswith('.edf'):
                drive_by_name[f['name']] = f
                if local is None:
                    edfs[f['name'][:-4]].append(dict(f, path=''))
    if local is not None:
        for folder in local['folders']:
            for f in folder['files']:
                if f['name'].lower().endswith('.edf'):
                    remote = drive_by_name.get(f['name'], {})
                    edfs[f['name'][:-4]].append({'id':remote.get('id',''), 'name':f['name'],
                                               'path':f['path'], 'size_bytes':f['bytes']})

    best = {}
    for _, r in excel_rows(src / 'stages_best_asq.xlsx'):
        if text(r['Use?']).lower() == 'yes':
            sid = text(r['subject_code'])
            if sid in best:
                raise ValueError(f'Multiple contributor-approved ASQ rows for {sid}')
            best[sid] = int(r['survey_id'])
    asq = defaultdict(dict)
    with zipfile.ZipFile(src / 'stages_asq.zip') as z:
        for member in sorted(z.namelist()):
            if member.endswith('.xlsx'):
                for _, r in excel_rows(BytesIO(z.read(member))):
                    sid = text(r['subject_code'])
                    if best.get(sid) == int(r['survey_id']):
                        asq[sid].update(r)

    cnb = defaultdict(list)
    validation_counts = Counter()
    with zipfile.ZipFile(src / 'stages_cnb.zip') as z:
        for member in sorted(z.namelist()):
            if not member.endswith('.xlsx'):
                continue
            mode = 'inlab' if 'inlab' in member else 'webservice'
            for rownum, r in excel_rows(BytesIO(z.read(member)), max_col=260):
                sid = text(r['test_sessions.subid'])
                score = num(r.get('PCET_A.PCET_ACC2'))
                cats, correct, incorrect = [num(r.get(k)) for k in ['PCET_A.PCET_CAT', 'PCET_A.PCET_CR', 'PCET_A.PCET_ER']]
                recomputed = ((cats + 1) * correct / (correct + incorrect)
                              if None not in (cats, correct, incorrect) and correct + incorrect > 0 else None)
                entry = {
                    'mode': mode, 'source_member': member, 'source_row': rownum,
                    'dataset_id': text(r.get('test_sessions.datasetid')),
                    'session': num(r.get('test_sessions_v.session')),
                    'date': day(r.get('modified.test_sessions_v.dotest')),
                    'battery_complete': num(r.get('test_sessions_v.battery_complete')),
                    'battery': text(r.get('test_sessions_v.battery')),
                    'status': text(r.get('PCET_A.status')),
                    'valid_code': text(r.get('PCET_A.valid_code')),
                    'skipped': truth(r.get('PCET_A.skipped')),
                    'pcet_accuracy2': score, 'pcet_recomputed': recomputed,
                    'pcet_recompute_difference': score - recomputed if score is not None and recomputed is not None else None,
                    'face_memory_correct': num(r.get('CPF_A.CPF_CR')),
                    'face_memory_status': text(r.get('CPF_A.status')),
                    'face_memory_validation': text(r.get('CPF_A.valid_code')),
                    'pvt_mean_reciprocal_rt': num(r.get('PVTB.PVTB_MEAN_RRT')),
                    'pvt_status': text(r.get('PVTB.status')),
                    'pvt_validation': text(r.get('PVTB.valid_code')),
                }
                cnb[sid].append(entry)
                validation_counts[(mode, entry['status'], entry['valid_code'])] += 1

    uploads = defaultdict(list)
    for r in pd.read_csv(src / 'stages_psg_upload.csv', dtype=str).fillna('').to_dict('records'):
        uploads[text(r['subject_id'])].append(r)
    srbd = defaultdict(list)
    for _, r in excel_rows(src / 'stages_srbd.xlsx'):
        srbd[text(r['s_code'])].append(r)
    education = defaultdict(list)
    for _, r in excel_rows(src / 'stages_siss.xlsx', sheet='STAGES_siss_mid 2020-06-28'):
        education[text(r['subject_id'])].append(r)

    all_ids = sorted(set(best) | set(cnb))
    attrition, eligible = [], []
    for sid in all_ids:
        a = asq.get(sid, {})
        reasons = []
        site, center = sid[:4], cfg['center_mapping'].get(sid[:4])
        if not center:
            reasons.append('unmapped_or_unavailable_center')
        if sid in cfg['excluded_duplicate_subjects']:
            reasons.append('known_duplicate_EDF_subject')
        if not a:
            reasons.append('no_contributor_approved_ASQ')
        # Different Drive copies with identical names must not be selected silently.
        if len(edfs.get(sid, [])) != 1:
            reasons.append('EDF_not_listed_or_ambiguous')
        rows = uploads.get(sid, [])
        psg_dates = {day(r['modified.date_psg']) for r in rows if day(r['modified.date_psg']) is not None}
        diagnostic = [r for r in rows if truth(r['type_psg_diagnostic'])
                      and not any(truth(r[k]) for k in ['type_psg_split_night','type_psg_mslt','type_psg_mwt','type_psg_fu_titration'])
                      and text(r['study_met_incl_excl_crit']).lower() == 'yes']
        if len(psg_dates) != 1:
            reasons.append('missing_or_multiple_PSG_dates')
        if not diagnostic:
            reasons.append('no_qualifying_diagnostic_PSG')
        psg_date = next(iter(psg_dates)) if len(psg_dates) == 1 else None
        sr = srbd.get(sid, [])
        def unique_sr(key):
            values = {text(r[key]) for r in sr if r.get(key) is not None}
            return next(iter(values)) if len(values) == 1 else None
        age_asq, age_psg = num(a.get('modified.dem_0110')), num(unique_sr('age'))
        sex_asq, sex_psg = sex_code(a.get('dem_0500')), sex_code(unique_sr('sex'))
        if age_asq is not None and age_psg is not None and abs(age_asq-age_psg) > 2:
            reasons.append('ASQ_PSG_age_conflict')
        if sex_asq is not None and sex_psg is not None and sex_asq != sex_psg:
            reasons.append('ASQ_PSG_sex_conflict')
        age = age_asq if age_asq is not None else age_psg
        sex = sex_asq if sex_asq is not None else sex_psg
        if age is None or age < cfg['primary_population']['minimum_age']:
            reasons.append('age_missing_or_under18')
        if sex is None:
            reasons.append('sex_missing')
        candidates = [r for r in cnb.get(sid, []) if r['mode'] == 'inlab'
                      and r['battery_complete'] == 1 and r['status'] == 'C1'
                      and r['valid_code'] == 'V' and not r['skipped']
                      and r['pcet_accuracy2'] is not None and 0 <= r['pcet_accuracy2'] <= 4]
        if not candidates:
            reasons.append('no_completed_valid_inlab_PCET')
        timed = [r for r in candidates if r['date'] is not None and psg_date is not None
                 and abs((r['date']-psg_date).days) <= cfg['primary_population']['max_absolute_psg_cnb_calendar_days']]
        if candidates and not timed:
            reasons.append('PCET_PSG_date_window_not_met')
        attrition.append({'subject_id': sid, 'site_prefix': site, 'center': center,
                          'eligible_before_signal_QC': not reasons, 'exclusion_reasons': '|'.join(reasons)})
        if reasons:
            continue
        cn = min(timed, key=lambda r: (r['date'], r['session'] or 999, r['dataset_id'], r['source_row']))
        if cn['pcet_recompute_difference'] is None or abs(cn['pcet_recompute_difference']) > 0.0002:
            raise ValueError(f'PCET score/count disagreement for {sid}')
        ed = education.get(sid, [])
        ed = sorted(ed, key=lambda r: (abs((day(r['modified.date_of_evaluation'])-psg_date).days)
                                      if day(r['modified.date_of_evaluation']) else 999999,
                                      text(r.get('high_educ_level_subject'))))
        educ = text(ed[0].get('high_educ_level_subject')) if ed else ''
        e = edfs[sid][0]
        bmi = num(a.get('dem_0800'))
        if bmi is None:
            bmi = num(unique_sr('bmi'))
        eligible.append({
            'subject_id': sid, 'site_prefix': site, 'center': center,
            'age': age, 'sex': sex, 'education': educ, 'bmi': bmi,
            'ahi': num(unique_sr('ahi')),
            'ess': total([a.get(f'ess_{i:02d}00') for i in range(1,9)]),
            'phq9': total([a.get(f'phq_{i:02d}00') for i in range(1,10)]),
            'asq_survey_id': best[sid], 'psg_date': psg_date.isoformat(),
            'cnb_date': cn['date'].isoformat(), 'cnb_psg_days': (cn['date']-psg_date).days,
            'cnb_dataset_id': cn['dataset_id'], 'cnb_session': cn['session'],
            'cnb_status': cn['status'], 'cnb_validation_code': cn['valid_code'],
            'pcet_accuracy2': cn['pcet_accuracy2'], 'pcet_recomputed': cn['pcet_recomputed'],
            'face_memory_correct': cn['face_memory_correct'],
            'face_memory_status': cn['face_memory_status'], 'face_memory_validation': cn['face_memory_validation'],
            'pvt_mean_reciprocal_rt': cn['pvt_mean_reciprocal_rt'],
            'pvt_status': cn['pvt_status'], 'pvt_validation': cn['pvt_validation'],
            'edf_drive_id': e['id'], 'edf_name': e['name'], 'edf_path': e['path'], 'edf_bytes': e['size_bytes'],
            'protocol_sha256': protocol_sha,
        })
    frame = pd.DataFrame(eligible).sort_values(['center','subject_id']).reset_index(drop=True)
    if frame.empty or not frame.subject_id.is_unique:
        raise ValueError('Empty or non-unique cohort')
    frame.to_csv(out / 'cohort_before_signal_qc.csv', index=False)
    pd.DataFrame(attrition).to_csv(out / 'attrition_by_subject.csv', index=False)
    extraction_cols = ['subject_id','site_prefix','center','age','sex','edf_drive_id','edf_name','edf_path','edf_bytes','protocol_sha256']
    frame[extraction_cols].to_csv(out / 'extraction_queue.csv', index=False)
    # Outcome-blind sampling of one subject per center for the initial signal/model pilot.
    pilot = frame[extraction_cols].copy()
    pilot['pilot_order'] = pilot.subject_id.map(lambda s: hashlib.sha256(f"{cfg['pilot_seed']}:{s}".encode()).hexdigest())
    pilot = pilot.sort_values(['center','pilot_order']).groupby('center', sort=True).head(1)
    pilot.drop(columns='pilot_order').to_csv(out / 'pilot_queue.csv', index=False)
    reason_counts = Counter(reason for r in attrition for reason in r['exclusion_reasons'].split('|') if reason)
    summary = {
        'protocol_sha256': protocol_sha, 'source_manifest_sha256': sha(src/'SHA256SUMS.json'),
        'inventory_sha256': sha(out/'psg_drive_inventory.json'),
        'inventory_complete': local is not None or all(x['listing_complete'] for x in inventory['folders']),
        'local_inventory_sha256': sha(local_inventory) if local is not None else None,
        'incomplete_inventory_folders': [] if local is not None else [x['site'] for x in inventory['folders'] if not x['listing_complete']],
        'contributor_approved_ASQ_subjects': len(best), 'CNB_unique_subject_strings': len(cnb),
        'candidate_subject_strings': len(all_ids), 'eligible_before_signal_QC': len(frame),
        'by_center': frame.groupby('center').size().to_dict(),
        'by_site_prefix': frame.groupby('site_prefix').size().to_dict(),
        'nonexclusive_exclusion_counts': dict(reason_counts),
        'covariate_missing': frame[['age','sex','education','bmi','ahi','ess','phq9']].replace('', None).isna().sum().to_dict(),
        'CNB_status_validation_counts': [{'mode': k[0], 'status': k[1], 'valid_code': k[2], 'n': v}
                                         for k,v in sorted(validation_counts.items())],
        'cohort_sha256': sha(out/'cohort_before_signal_qc.csv'),
        'extraction_queue_sha256': sha(out/'extraction_queue.csv'),
        'status': 'cohort_built_no_EEG_associations_computed',
    }
    (base/'results/cohort_build.json').write_text(json.dumps(summary, indent=2)+'\n')
    print(json.dumps({k: summary[k] for k in ['eligible_before_signal_QC','by_center','covariate_missing','nonexclusive_exclusion_counts']}, indent=2))


if __name__ == '__main__':
    main()
