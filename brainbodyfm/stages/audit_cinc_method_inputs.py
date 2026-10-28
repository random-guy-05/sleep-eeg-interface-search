"""Audit input identity, scoring, checkpoint projection and historical reproduction.

Never changes a result, label definition, score direction, or existing cache.
Only aggregate audit evidence is written to the requested output.
"""
from __future__ import annotations
import argparse
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime,timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import sys
import zipfile
import numpy as np
import pandas as pd
import torch
from openpyxl import load_workbook
from threadpoolctl import threadpool_limits
import stable4_extension
from stable4_extension import load_inputs
from analyze import STABLE4
import cinc_method_replication as replication


def main():
    ap=argparse.ArgumentParser();ap.add_argument('--data-root',type=Path,required=True);ap.add_argument('--entry-source',type=Path,required=True);ap.add_argument('--checkpoint',type=Path,required=True);ap.add_argument('--output',type=Path,required=True);a=ap.parse_args()
    root=a.data_root.resolve();project=root.parents[1]
    entry=replication.load_entry(a.entry_source)
    stable4_extension.BASE=root;stable4_extension.OUT=root/'results/stable4_extension'
    frame,semantic,latent,bhs,names=load_inputs()
    selected=[names.index(h) for h in STABLE4]
    report=dict(created_utc=datetime.now(timezone.utc).isoformat(),entry_commit=entry['commit'],checks={})
    checks=report['checks']
    checks['input_dimensions']=dict(n=len(frame),unique_subjects=int(frame.subject_id.nunique()),semantic=list(semantic.shape),latent=list(latent.shape),finite=bool(np.isfinite(semantic).all() and np.isfinite(latent).all() and np.isfinite(bhs).all()))
    max_heads=max_latent=max_bhs=0.;versions=Counter();checkpoints=Counter();head_orders=Counter()
    def inspect_subject(item):
        i,sid=item
        p=root/'cache/assumed_uv/subjects'/f'{sid}.npz'
        with np.load(p,allow_pickle=False) as z:
            errors=(float(np.max(np.abs(z['semantic']-semantic[i]))),
                    float(np.max(np.abs(z['latent']-latent[i]))),
                    abs(float(z['brain_health_score'])-bhs[i]),
                    z['semantic_names'].tolist()==names)
        meta=json.loads(p.with_suffix('.json').read_text())
        if hashlib.sha256(p.read_bytes()).hexdigest()!=meta['payload_sha256']:
            raise ValueError('Cache hash mismatch')
        return errors,meta
    with ThreadPoolExecutor(max_workers=8) as pool:
        for n,(errors,meta) in enumerate(pool.map(inspect_subject,enumerate(frame.subject_id)),1):
            h,l,b,o=errors
            max_heads=max(max_heads,h);max_latent=max(max_latent,l);max_bhs=max(max_bhs,b);head_orders[o]+=1
            versions[json.dumps(meta['source_hashes'],sort_keys=True)]+=1
            checkpoints[meta['checkpoint_sha256']]+=1
            if n%50==0:print('subject caches checked',n,flush=True)
    checks['every_subject_cache_matches_joined_input']=dict(checked=len(frame),semantic_max_abs_difference=max_heads,latent_max_abs_difference=max_latent,bhs_max_abs_difference=max_bhs,matching_head_order=head_orders[True],checkpoint_counts=dict(checkpoints),extraction_source_versions=[dict(hashes=json.loads(k),n=v) for k,v in versions.items()])
    print('cache identity audit complete',flush=True)
    # Select only the four needed head tensors from a memory-mapped checkpoint.
    # This independent projection does not call the STAGES head-export function.
    checkpoint=torch.load(a.checkpoint,map_location='cpu',weights_only=False,mmap=True)
    state=checkpoint['state_dict'];projected=[]
    for name in STABLE4:
        suffix=name.removeprefix('pred_regression_')
        w=state['heads.final_head_'+suffix+'.weight'].detach().numpy().astype(np.float64)
        b=state['heads.final_head_'+suffix+'.bias'].detach().numpy().astype(np.float64)
        projected.append((latent.astype(np.float64)@w.T+b).reshape(-1))
    projected=np.column_stack(projected)
    error=np.abs(projected-semantic[:,selected]);scale=np.maximum(np.std(semantic[:,selected],axis=0),1e-12)
    checks['independent_checkpoint_projection']=dict(heads=STABLE4,max_absolute_error_by_head=error.max(axis=0).tolist(),max_error_in_head_sd_units=(error.max(axis=0)/scale).tolist(),correlations=[float(np.corrcoef(projected[:,j],semantic[:,selected[j]])[0,1]) for j in range(4)])
    print('checkpoint projection audit complete',flush=True)
    # Independently retrieve the selected original cognitive rows.
    wanted={str(r.subject_id):(str(r.cnb_dataset_id),float(r.cnb_session)) for r in frame.itertuples()}
    original={};duplicate=0
    with zipfile.ZipFile(root/'data/source/stages_cnb.zip') as archive:
        member=next(n for n in archive.namelist() if 'inlab' in n and n.endswith('.xlsx'))
        book=load_workbook(BytesIO(archive.read(member)),read_only=True,data_only=True)
        sheet=book.active;columns=next(sheet.iter_rows(max_row=1,values_only=True))
        for values in sheet.iter_rows(min_row=2,values_only=True):
            r=dict(zip(columns,values));sid=str(r.get('test_sessions.subid','')).strip()
            if sid not in wanted:continue
            dataset=str(r.get('test_sessions.datasetid','')).strip();session=r.get('test_sessions_v.session')
            if session is None or (dataset,float(session))!=wanted[sid]:continue
            duplicate+=sid in original
            original[sid]=r
        book.close()
    missing=set(frame.subject_id)-set(original)
    if missing:raise ValueError('Selected cognitive source rows not recovered: '+str(len(missing)))
    max_pcet=max_formula=max_memory=0.;memory_missing=0;bad_status=0
    for r in frame.itertuples():
        s=original[r.subject_id];pcet=float(s['PCET_A.PCET_ACC2'])
        cr=float(s['PCET_A.PCET_CR']);er=float(s['PCET_A.PCET_ER']);cat=float(s['PCET_A.PCET_CAT'])
        max_pcet=max(max_pcet,abs(pcet-r.pcet_accuracy2))
        max_formula=max(max_formula,abs(pcet-(cat+1)*cr/(cr+er)))
        memory=s.get('CPF_A.CPF_CR')
        if memory is None or pd.isna(r.face_memory_correct):
            memory_missing+=1
            if not(memory is None and pd.isna(r.face_memory_correct)):raise ValueError('Memory missingness mismatch')
        else:max_memory=max(max_memory,abs(float(memory)-r.face_memory_correct))
        bad_status+=(str(s['PCET_A.status'])!='C1' or str(s['PCET_A.valid_code'])!='V')
    checks['original_cognitive_rows']=dict(checked=len(frame),duplicate_matching_rows=duplicate,missing_matches=len(missing),pcet_max_abs_difference=max_pcet,pcet_formula_max_abs_rounding_error=max_formula,face_memory_max_abs_difference=max_memory,face_memory_missing=memory_missing,invalid_pcet_status=bad_status,score_direction='higher accuracy and more correct face-memory responses mean better performance')
    print('original cognitive source-row audit complete',flush=True)
    # Verify the same scorer still reproduces the known HSP development result.
    old=pd.read_csv(project/'experiments/philosophers_stone/pred_heads_full_338.csv')
    oldscores=replication.ENTRY._site_loso_stable4_oof(old[STABLE4].to_numpy(float),old.sex.to_numpy(float),old.site_id.to_numpy(str))
    site_values={};pairs={}
    for site in sorted(old.site_id.unique()):
        use=old.site_id.eq(site).to_numpy()
        v,n=replication.ENTRY.pair_auc(oldscores[use],old.target.to_numpy(int)[use],old.age.to_numpy(float)[use],2.)
        site_values[site]=v;pairs[site]=n
    checks['historical_HSP_development_reproduction']=dict(n=len(old),site_auc=site_values,eligible_pairs=pairs,macro_auc=float(np.mean(list(site_values.values()))),within_site_pooled_auc=float(np.average(list(site_values.values()),weights=list(pairs.values()))))
    oldlatent=pd.read_csv(project/'brainbodyfm/cache/raw_latent_1024_338.csv').set_index('subject_id').loc[old.subject_id].to_numpy(float)
    oldproj=[]
    for name in STABLE4:
        suffix=name.removeprefix('pred_regression_');w=state['heads.final_head_'+suffix+'.weight'].detach().numpy().astype(float);b=state['heads.final_head_'+suffix+'.bias'].detach().numpy().astype(float)
        oldproj.append((oldlatent@w.T+b).reshape(-1))
    oldproj=np.column_stack(oldproj);oldheads=old[STABLE4].to_numpy(float)
    checks['historical_latents_projected_with_current_checkpoint']=dict(max_abs_error_by_head=np.max(np.abs(oldproj-oldheads),axis=0).tolist(),correlations=[float(np.corrcoef(oldproj[:,j],oldheads[:,j])[0,1]) for j in range(4)],note='Historical extraction checkpoint identity remains unverified; this checks compatibility of head tensors with recovered latents.')
    # Independently score each saved STAGES prediction using the submitted metric.
    summaries=json.loads((root/'results/cinc_method_replication/summary.json').read_text())
    metricchecks={}
    framebyid=frame.set_index('subject_id')
    for task,result in summaries['analyses'].items():
        p=np.load(root/'results/cinc_method_replication'/(task+'_private_predictions.npz'))
        aligned=framebyid.loc[p['ids']];values=[];ns=[]
        for center in result['counts']:
            ix=aligned.center.eq(center).to_numpy()
            v,n=replication.ENTRY.pair_auc(p['scores'][ix,0],p['labels'][ix],aligned.age.to_numpy(float)[ix],2.)
            values.append(v);ns.append(n)
        actual=float(np.average(values,weights=ns));expected=result['age_conditioned']['stable4']
        np.testing.assert_allclose(actual,expected,atol=1e-14,rtol=0)
        metricchecks[task]=dict(submitted_metric_value=actual,absolute_difference=abs(actual-expected),eligible_pairs=sum(ns),matches=True)
    checks['independent_submitted_metric']=metricchecks
    report['status']='complete';a.output.parent.mkdir(parents=True,exist_ok=True)
    a.output.write_text(json.dumps(report,indent=2,allow_nan=False)+'\n')
    print(json.dumps(report,indent=2),flush=True)


if __name__=='__main__':
    torch.set_num_threads(1)
    with threadpool_limits(limits=1):main()
